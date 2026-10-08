"""Conversion Beamer → PPTX pour la plateforme des JFR (PowerPoint 2019, 16:9, < 300 Mo).

Chaque page du PDF est rendue en PNG (pdftoppm, 300 dpi) et posée plein cadre sur une diapo
vierge ; aucun texte PowerPoint, donc aucune police à embarquer. Sur la dernière diapo, la vidéo
de la boucle de scopie (figures/fin_bruit_qr.mp4) remplace l'image : lecture automatique, en
boucle, plein cadre, l'image fixe de la page servant d'affiche (poster frame) et de secours.

Usage : python make_pptx.py [PDF] [SORTIE.pptx]
Prérequis : pdftoppm (poppler), python-pptx ; la vidéo H.264 est embarquée, pas liée.
"""
import subprocess
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

racine = Path(__file__).resolve().parent
pdf = Path(sys.argv[1]) if len(sys.argv) > 1 else racine / "build" / "jfr2026.pdf"
out = Path(sys.argv[2]) if len(sys.argv) > 2 else racine / "build" / "JFR2026_Ammour_metriques_RI.pptx"
video = racine / "figures" / "fin_bruit_qr.mp4"
png_dir = racine / "build" / "pptx_png"
DPI = 300

png_dir.mkdir(parents=True, exist_ok=True)
for f in png_dir.glob("p-*.png"):
    f.unlink()
subprocess.run(["pdftoppm", "-png", "-r", str(DPI), str(pdf), str(png_dir / "p")], check=True)
pages = sorted(png_dir.glob("p-*.png"))
assert pages, "aucune page rendue"

def autoplay_loop(slide, spid):
    """Remplace le minutage écrit par python-pptx (lecture au clic) par celui que PowerPoint
    produit pour « Démarrer : automatiquement » + « Lire en boucle jusqu'à l'arrêt »."""
    from lxml import etree
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    xml = f"""<p:timing xmlns:p="{ns}"><p:tnLst><p:par>
<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="indefinite" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0" repeatCount="indefinite"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst><p:cond evt="onEnd" delay="0"><p:tn val="7"/></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""
    el = slide._element
    for old in el.findall(f"{{{ns}}}timing"):
        el.remove(old)
    el.append(etree.fromstring(xml))


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)      # 16:9, 33,87 × 19,05 cm
vierge = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
for i, png in enumerate(pages, 1):
    slide = prs.slides.add_slide(vierge)
    if i == len(pages) and video.exists():
        film = slide.shapes.add_movie(str(video), 0, 0, W, H, poster_frame_image=str(png), mime_type="video/mp4")
        autoplay_loop(slide, film.shape_id)
    else:
        slide.shapes.add_picture(str(png), 0, 0, W, H)

prs.save(out)
print(f"{len(pages)} diapos → {out} ({out.stat().st_size / 1e6:.1f} Mo) ; vidéo sur la dernière : {video.exists()}")
